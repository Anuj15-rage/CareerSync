"""
Stage 4.3 Acceptance Tests — Frontend-Backend Integration & CORS Contract
Updated for Phase 1 token security and isolated SQLite execution.
"""
import io
import sys
import uuid
from pathlib import Path

# Add project root to sys.path so backend imports resolve cleanly
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import docx
from fastapi.testclient import TestClient

from backend.main import app
import backend.services.rate_limiter as rl
from tests.test_sqlite_fixtures import setup_sqlite_test_env

ALLOWED_TEST_ORIGIN = "http://127.0.0.1:5500"
DISALLOWED_ORIGIN = "http://unauthorized-domain.com"


def reset_rate_limiter():
    rl._UPLOAD_HISTORY.clear()


def create_sample_docx() -> bytes:
    """Create in-memory DOCX resume with realistic profile data."""
    doc = docx.Document()
    doc.add_paragraph("Sarah Connor")
    doc.add_paragraph("Email: sarah.connor@example.com | Phone: +1-555-019-2834 | Austin, TX")
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph(
        "University of Texas at Austin\n"
        "Bachelor of Science in Computer Science | 2017 - 2021\n"
        "CGPA: 3.85 / 4.0"
    )
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph(
        "Languages: Python, TypeScript, SQL, HTML5, CSS3, C++\n"
        "Frameworks: FastAPI, React, Node.js\n"
        "Tools & Cloud: Docker, Git, GitHub, AWS, CI/CD Pipelines"
    )
    doc.add_heading("WORK EXPERIENCE", level=1)
    doc.add_paragraph(
        "Full Stack Developer\n"
        "Cyberdyne Innovations | Jun 2021 - Present\n"
        "- Developed microservices using FastAPI and PostgreSQL\n"
        "- Maintained containerized deployments using Docker and AWS"
    )
    doc.add_heading("PROJECTS", level=1)
    doc.add_paragraph(
        "Automated Health Monitor\n"
        "Technologies: Python, FastAPI, Docker, React\n"
        "- Real-time telemetry dashboard with RESTful APIs"
    )

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_cors_preflight_and_headers(client):
    """Verify CORS preflight OPTIONS requests return valid headers for allowed local origins."""
    headers = {
        "Origin": ALLOWED_TEST_ORIGIN,
        "Access-Control-Request-Method": "POST",
        "Access-Control-Request-Headers": "Content-Type",
    }
    response = client.options("/api/resumes/upload", headers=headers)
    assert response.status_code == 200, f"Expected 200 on OPTIONS preflight, got {response.status_code}"
    assert response.headers.get("access-control-allow-origin") == ALLOWED_TEST_ORIGIN
    assert "POST" in response.headers.get("access-control-allow-methods", "")

    # Disallowed origin should not get access-control-allow-origin header
    bad_headers = {
        "Origin": DISALLOWED_ORIGIN,
        "Access-Control-Request-Method": "POST",
    }
    bad_response = client.options("/api/resumes/upload", headers=bad_headers)
    assert bad_response.headers.get("access-control-allow-origin") is None or bad_response.headers.get("access-control-allow-origin") != DISALLOWED_ORIGIN

    print("  [PASS] test_cors_preflight_and_headers: CORS headers properly enforced.")


def test_frontend_upload_and_analyze_contract(client):
    """
    Verify complete upload -> analyze workflow with CORS headers
    and validate the JSON contract expected by frontend/script.js.
    """
    reset_rate_limiter()
    docx_bytes = create_sample_docx()
    unique_name = f"sarah_connor_{uuid.uuid4().hex[:6]}.docx"

    # Step 1: Upload with Origin header simulating frontend
    files = {"file": (unique_name, docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    headers = {"Origin": ALLOWED_TEST_ORIGIN}
    upload_res = client.post("/api/resumes/upload", files=files, headers=headers)

    assert upload_res.status_code == 200, f"Upload failed: {upload_res.text}"
    assert upload_res.headers.get("access-control-allow-origin") == ALLOWED_TEST_ORIGIN
    upload_data = upload_res.json()

    assert "id" in upload_data
    assert "access_token" in upload_data
    resume_id = upload_data["id"]
    token = upload_data["access_token"]

    assert upload_data["parsing_status"] == "PARSED"
    assert upload_data["file_type"] == "DOCX"
    assert upload_data["word_count"] > 0

    # Step 2: Trigger Analysis with Origin and X-Resume-Token headers
    analyze_headers = {
        "Origin": ALLOWED_TEST_ORIGIN,
        "X-Resume-Token": token,
    }
    analyze_res = client.post(f"/api/resumes/{resume_id}/analyze", headers=analyze_headers)
    assert analyze_res.status_code == 200, f"Analysis failed: {analyze_res.text}"
    assert analyze_res.headers.get("access-control-allow-origin") == ALLOWED_TEST_ORIGIN
    analysis_data = analyze_res.json()

    # Step 3: Validate exact frontend contract fields
    cand = analysis_data.get("candidate", {})
    assert cand.get("candidate_name") == "Sarah Connor"
    assert cand.get("candidate_email") == "sarah.connor@example.com"
    assert cand.get("candidate_phone") == "+1-555-019-2834"

    # Skills array
    skills = analysis_data.get("skills", [])
    assert len(skills) >= 6
    for s in skills:
        assert "skill_id" in s
        assert "skill_name" in s
        assert "category" in s
        assert "extraction_source" in s
        assert s["extraction_source"] in ("SKILLS_SECTION", "PROJECTS", "EXPERIENCE", "EDUCATION")

    # Education array
    edu = analysis_data.get("education", [])
    assert len(edu) >= 1
    assert "University of Texas" in edu[0].get("institution", "")
    assert "Bachelor of Science" in edu[0].get("degree", "")

    # Projects array
    projects = analysis_data.get("projects", [])
    assert len(projects) >= 1
    assert "Automated Health Monitor" in projects[0].get("title", "")
    assert "Python" in projects[0].get("technologies_used", "")

    # Experience array
    experience = analysis_data.get("experience", [])
    assert len(experience) >= 1
    assert "Cyberdyne Innovations" in experience[0].get("organization", "")
    assert "Full Stack Developer" in experience[0].get("role_title", "")

    print(f"  [PASS] test_frontend_upload_and_analyze_contract: JSON schema matches frontend requirements (id={resume_id}).")


def run_all_tests():
    print("\n=== Running Stage 4.3 Acceptance Tests (Frontend-Backend Integration) ===")
    setup_sqlite_test_env(app)
    client = TestClient(app)

    test_cors_preflight_and_headers(client)
    test_frontend_upload_and_analyze_contract(client)

    print("\n=== All Stage 4.3 Acceptance Tests Passed Successfully! ===\n")


if __name__ == "__main__":
    run_all_tests()
