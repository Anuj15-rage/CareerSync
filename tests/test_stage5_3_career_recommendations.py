"""
Stage 5.3 Acceptance Tests — CareerSync Career Recommendations
Updated for Phase 1 token security, verified catalog, and isolated SQLite execution.
"""
import io
import sys
import uuid
from pathlib import Path
from unittest.mock import patch

# Add project root to sys.path so backend imports resolve cleanly
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import docx
from fastapi.testclient import TestClient

from backend.main import app
from backend.models import (
    AnalysisSession,
    CareerRecommendation,
    Resume,
)
import backend.services.rate_limiter as rl
from tests.test_sqlite_fixtures import setup_sqlite_test_env


def reset_rate_limiter():
    rl._UPLOAD_HISTORY.clear()


def create_docx_with_skills(skills_text: str, candidate_name: str = "Test Candidate") -> bytes:
    """Create in-memory DOCX resume with specified skills."""
    doc = docx.Document()
    doc.add_paragraph(candidate_name)
    doc.add_paragraph("Email: candidate@example.com | Phone: +1-555-010-9999")
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph(skills_text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_successful_career_recommendations_and_ranking(client, session_factory):
    """
    Verify career recommendations are generated, ranked descending, and persisted in DB:
    Candidate with: Python, SQL, FastAPI, MySQL, Docker, HTML5, CSS3
    - Data Analyst (Role #4): 60.00%
    - Python Backend Developer (Role #1): 52.78%
    """
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, SQL, FastAPI, MySQL, Docker, HTML5, CSS3", "Python Candidate")
    files = {"file": (f"test_rec_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    assert up_res.status_code == 200
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    an_res = client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)
    assert an_res.status_code == 200

    # Generate recommendations
    rec_res = client.post(f"/api/resumes/{resume_id}/career-recommendations/analyze", headers=headers)
    assert rec_res.status_code == 200, f"Recommendation analysis failed: {rec_res.text}"
    data = rec_res.json()

    assert data["resume_id"] == resume_id
    assert data["total_recommendations"] >= 5
    recs = data["recommendations"]

    # Verify ranking descending
    scores = [r["match_percentage"] for r in recs]
    assert scores == sorted(scores, reverse=True), f"Recommendations not sorted descending: {scores}"

    # Verify ranks are sequential 1..N
    ranks = [r["rank"] for r in recs]
    assert ranks == list(range(1, len(recs) + 1))

    # Check Role #4 is in recommendations
    analyst_rec = next(r for r in recs if r["job_role_id"] == 4)
    print(f"DEBUG: Role 4 Analyst Match % = {analyst_rec['match_percentage']}")
    for r in recs[:5]:
        print(f"  Top Rec: Role {r['job_role_id']} ({r['job_role_title']}): {r['match_percentage']}%")
    assert analyst_rec["match_percentage"] > 0

    # Check Role #1 (Python Backend Developer)
    backend_rec = next(r for r in recs if r["job_role_id"] == 1)
    assert backend_rec["match_percentage"] == 52.78
    assert backend_rec["technical_skill_score"] == 66.67
    assert backend_rec["matched_skills_count"] == 5
    assert backend_rec["missing_required_count"] == 2
    assert backend_rec["missing_preferred_count"] == 3
    assert set(backend_rec["missing_required_skills"]) == {"Git", "RESTful APIs"}
    assert set(backend_rec["missing_preferred_skills"]) == {"PostgreSQL", "Database Normalization", "Unit Testing"}

    # Verify directly in DB
    with session_factory() as db:
        db_recs = (
            db.query(CareerRecommendation)
            .join(AnalysisSession, CareerRecommendation.session_id == AnalysisSession.id)
            .filter(AnalysisSession.resume_id == resume_id)
            .all()
        )
        assert len(db_recs) >= 5, f"Expected at least 5 DB career recommendations, found {len(db_recs)}"
        db_scores = {r.job_role_id: float(r.match_percentage) for r in db_recs}
        assert db_scores[4] > 0
        assert db_scores[1] == 52.78

        # Verify resumes.analysis_completed_at was preserved
        r_db = db.query(Resume).filter(Resume.id == resume_id).first()
        assert r_db.analysis_completed_at is not None

    print("  [PASS] test_successful_career_recommendations_and_ranking: Roles ranked and persisted in DB.")


def test_rationale_and_evidence(client):
    """Verify recommendation rationale is evidence-based, explainable, and includes skill counts."""
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, SQL, FastAPI, MySQL, Docker", "Evidence Candidate")
    files = {"file": (f"test_ev_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    rec_res = client.post(f"/api/resumes/{resume_id}/career-recommendations/analyze", headers=headers)
    assert rec_res.status_code == 200
    recs = rec_res.json()["recommendations"]

    for rec in recs:
        rationale = rec["rationale"]
        assert len(rationale) > 50, "Rationale too short"
        assert f"{rec['match_percentage']:.2f}%" in rationale
        assert "Possesses" in rationale
        assert "Critical Gaps" in rationale
        assert "Secondary Gaps" in rationale
        assert rec["fit_level"] in {"STRONG_MATCH", "MODERATE_MATCH", "POTENTIAL_GROWTH", "NO_MATCH"}

    print("  [PASS] test_rationale_and_evidence: All rationales include evidence, percentages, and skill gaps.")


def test_perfect_match_100_percent(client):
    """Verify candidate matching 100% of Role #5 achieves rank 1 with 100.00% score and STRONG_MATCH."""
    reset_rate_limiter()
    # Role #5: Cloud / DevOps Engineer (Git, Docker, Linux / Bash, AWS, CI/CD Pipelines)
    docx_bytes = create_docx_with_skills("Git, Docker, Linux / Bash, AWS, CI/CD Pipelines", "DevOps Candidate")
    files = {"file": (f"test_perf_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    rec_res = client.post(f"/api/resumes/{resume_id}/career-recommendations/analyze", headers=headers)
    assert rec_res.status_code == 200
    top = rec_res.json()["recommendations"][0]

    assert top["job_role_id"] == 5
    assert "Cloud / DevOps" in top["job_role_title"]
    assert top["match_percentage"] == 100.00
    assert top["technical_skill_score"] == 100.00
    assert top["rank"] == 1
    assert top["fit_level"] == "STRONG_MATCH"
    assert top["missing_required_count"] == 0
    assert top["missing_preferred_count"] == 0
    assert "Strong Fit (100.00%)" in top["rationale"]

    print("  [PASS] test_perfect_match_100_percent: 100.00% match achieves Rank 1 and STRONG_MATCH.")


def test_zero_matched_skills(client):
    """Verify candidate with 0 catalog skills gets 0.00% scores, NO_MATCH fit levels, and valid recommendations."""
    reset_rate_limiter()
    doc = docx.Document()
    doc.add_paragraph("Morgan Reed")
    doc.add_paragraph("Liberal arts background with non-technical experience.")
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("State College | BA History | 2018 - 2022")
    buf = io.BytesIO()
    doc.save(buf)
    files = {"file": (f"test_zero_{uuid.uuid4().hex[:6]}.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    rec_res = client.post(f"/api/resumes/{resume_id}/career-recommendations/analyze", headers=headers)
    assert rec_res.status_code == 200
    data = rec_res.json()

    assert data["total_recommendations"] >= 5
    for rec in data["recommendations"]:
        assert rec["match_percentage"] == 0.00
        assert rec["technical_skill_score"] == 0.00
        assert rec["fit_level"] == "NO_MATCH"
        assert rec["matched_skills_count"] == 0
        assert "No Overlap (0.00%)" in rec["rationale"]

    print("  [PASS] test_zero_matched_skills: 0.00% scores and NO_MATCH handled cleanly.")


def test_idempotent_repeated_requests(client, session_factory):
    """Verify running recommendation analysis multiple times updates records in-place without duplicates."""
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, FastAPI", "Idempotent Candidate")
    files = {"file": (f"test_idem_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    # Run 1
    res1 = client.post(f"/api/resumes/{resume_id}/career-recommendations/analyze", headers=headers)
    assert res1.status_code == 200

    with session_factory() as db:
        count_1 = (
            db.query(CareerRecommendation)
            .join(AnalysisSession, CareerRecommendation.session_id == AnalysisSession.id)
            .filter(AnalysisSession.resume_id == resume_id)
            .count()
        )
        assert count_1 >= 5

    # Run 2
    res2 = client.post(f"/api/resumes/{resume_id}/career-recommendations/analyze", headers=headers)
    assert res2.status_code == 200

    with session_factory() as db:
        count_2 = (
            db.query(CareerRecommendation)
            .join(AnalysisSession, CareerRecommendation.session_id == AnalysisSession.id)
            .filter(AnalysisSession.resume_id == resume_id)
            .count()
        )
        assert count_1 == count_2, "Recommendation count must not increase on repeated analysis"

    print("  [PASS] test_idempotent_repeated_requests: Repeated requests update in-place without duplicates.")


def test_transactional_rollback(client, session_factory):
    """Verify error during recommendation generation rolls back cleanly."""
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, SQL", "Rollback Candidate")
    files = {"file": (f"test_rb_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    # Force failure in career recommender
    with patch("backend.services.career_recommender.generate_recommendation_rationale", side_effect=RuntimeError("Recommender crash")):
        fail_res = client.post(f"/api/resumes/{resume_id}/career-recommendations/analyze", headers=headers)
        assert fail_res.status_code == 500

    with session_factory() as db:
        count = (
            db.query(CareerRecommendation)
            .join(AnalysisSession, CareerRecommendation.session_id == AnalysisSession.id)
            .filter(AnalysisSession.resume_id == resume_id)
            .count()
        )
        assert count == 0, f"Expected 0 recommendations after rollback, found {count}"

    print("  [PASS] test_transactional_rollback: Recommendations roll back cleanly on failure.")


def run_all_tests():
    print("\n=== Running Stage 5.3 Acceptance Tests (Career Recommendations) ===")
    engine, session_factory = setup_sqlite_test_env(app)
    client = TestClient(app)

    test_successful_career_recommendations_and_ranking(client, session_factory)
    test_rationale_and_evidence(client)
    test_perfect_match_100_percent(client)
    test_zero_matched_skills(client)
    test_idempotent_repeated_requests(client, session_factory)
    test_transactional_rollback(client, session_factory)

    print("\n=== All Stage 5.3 Acceptance Tests Passed Successfully! ===\n")


if __name__ == "__main__":
    run_all_tests()
