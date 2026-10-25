"""
Stage 4.1 Acceptance Tests — CareerSync Resume Upload and Text Extraction
Updated for Phase 1 token security, 5 MB upload limits, and isolated SQLite execution.
"""
import io
import sys
from pathlib import Path

# Add project root to sys.path so backend imports resolve cleanly
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import docx
from fastapi.testclient import TestClient
from pypdf import PdfReader, PdfWriter

from backend.main import app
from tests.test_sqlite_fixtures import setup_sqlite_test_env


def make_valid_pdf_bytes(text: str = "John Doe - Software Engineer with Python and SQL experience") -> bytes:
    """Generate in-memory valid PDF bytes containing readable text."""
    raw_pdf = f"""%PDF-1.4
1 0 obj
<< /Type /Catalog /Pages 2 0 R >>
endobj
2 0 obj
<< /Type /Pages /Kids [3 0 R] /Count 1 >>
endobj
3 0 obj
<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>
endobj
4 0 obj
<< /Length {len(text) + 30} >>
stream
BT
/F1 12 Tf
72 712 Td
({text}) Tj
ET
endstream
endobj
5 0 obj
<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>
endobj
xref
0 6
0000000000 65535 f 
0000000009 00000 n 
0000000058 00000 n 
0000000115 00000 n 
0000000244 00000 n 
0000000350 00000 n 
trailer
<< /Size 6 /Root 1 0 R >>
startxref
424
%%EOF""".encode("latin-1")
    reader = PdfReader(io.BytesIO(raw_pdf))
    writer = PdfWriter()
    for p in reader.pages:
        writer.add_page(p)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_blank_pdf_bytes() -> bytes:
    """Generate in-memory valid PDF bytes containing a blank page (simulating scanned/empty PDF)."""
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def make_valid_docx_bytes(paragraphs: list[str]) -> bytes:
    """Generate in-memory valid DOCX bytes containing body paragraphs and a table."""
    doc = docx.Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Proficiency"
    table.cell(1, 0).text = "Python"
    table.cell(1, 1).text = "Advanced"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def reset_rate_limiter():
    import backend.services.rate_limiter as rl
    rl._UPLOAD_HISTORY.clear()


def test_valid_pdf_upload(client):
    """TC-4.1-01: Upload valid PDF, verify 200 OK, text extraction, and access token."""
    reset_rate_limiter()
    pdf_bytes = make_valid_pdf_bytes("John Doe - Senior Backend Engineer")
    files = {"file": ("john_doe_resume.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/resumes/upload", files=files)

    assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
    data = response.json()
    assert "id" in data
    assert "access_token" in data
    assert len(data["access_token"]) >= 40
    assert data["file_type"] == "PDF"
    assert data["parsing_status"] == "PARSED"
    assert data["word_count"] > 0
    print(f"  [PASS] test_valid_pdf_upload: PDF uploaded successfully (ID: {data['id']}).")
    return data["id"], data["access_token"]


def test_valid_docx_upload(client):
    """TC-4.1-02: Upload valid DOCX, verify 200 OK, text extraction, and access token."""
    reset_rate_limiter()
    docx_bytes = make_valid_docx_bytes([
        "Jane Smith",
        "Full Stack Developer specializing in React, Node.js, and SQL.",
        "Education: B.S. in Computer Science",
    ])
    files = {
        "file": (
            "jane_smith_cv.docx",
            docx_bytes,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    response = client.post("/api/resumes/upload", files=files)

    assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
    data = response.json()
    assert "id" in data
    assert "access_token" in data
    assert data["file_type"] == "DOCX"
    assert data["parsing_status"] == "PARSED"
    assert data["word_count"] > 0
    print(f"  [PASS] test_valid_docx_upload: DOCX uploaded successfully (ID: {data['id']}).")
    return data["id"], data["access_token"]


def test_get_resume_metadata_endpoint(client, resume_id, access_token):
    """TC-4.1-03: Retrieve resume metadata with valid X-Resume-Token."""
    headers = {"X-Resume-Token": access_token}
    response = client.get(f"/api/resumes/{resume_id}", headers=headers)

    assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
    data = response.json()
    assert data["id"] == resume_id
    assert data["file_type"] == "PDF"
    assert data["parsing_status"] == "PARSED"
    print(f"  [PASS] test_get_resume_metadata_endpoint: Retrieved metadata for ID {resume_id} using token.")


def test_get_nonexistent_resume_metadata(client):
    """TC-4.1-04: Nonexistent resume returns 404 Not Found."""
    headers = {"X-Resume-Token": "nonexistent_token_placeholder"}
    response = client.get("/api/resumes/999999", headers=headers)
    assert response.status_code == 404, f"Expected 404, got {response.status_code}"
    print("  [PASS] test_get_nonexistent_resume_metadata: 404 returned for nonexistent resume.")


def test_invalid_extension_rejected(client):
    """TC-4.1-05: Unsupported extension (.txt) rejected with 400 Bad Request."""
    files = {"file": ("resume.txt", b"Some plain text resume", "text/plain")}
    response = client.post("/api/resumes/upload", files=files)
    assert response.status_code == 400, f"Expected 400, got {response.status_code}"
    print("  [PASS] test_invalid_extension_rejected: .txt rejected with 400.")


def test_spoofed_signature_rejected(client):
    """TC-4.1-06: File with .pdf extension but invalid signature rejected with 400."""
    fake_content = b"This is clearly not a real PDF file despite the extension"
    files = {"file": ("fake.pdf", fake_content, "application/pdf")}
    response = client.post("/api/resumes/upload", files=files)
    assert response.status_code == 400, f"Expected 400, got {response.status_code}"
    assert "signature mismatch" in response.json()["detail"].lower()
    print("  [PASS] test_spoofed_signature_rejected: Spoofed signature rejected with 400.")


def test_empty_zero_byte_file_rejected(client):
    """TC-4.1-07: Empty file (0 bytes) rejected with 400 Bad Request."""
    files = {"file": ("empty.pdf", b"", "application/pdf")}
    response = client.post("/api/resumes/upload", files=files)
    assert response.status_code == 400, f"Expected 400, got {response.status_code}"
    print("  [PASS] test_empty_zero_byte_file_rejected: Empty 0-byte file rejected with 400.")


def test_scanned_or_blank_pdf_graceful_handling(client):
    """TC-4.1-08: Blank or image-only PDF handled gracefully (200 OK with FAILED parsing status)."""
    reset_rate_limiter()
    blank_pdf_bytes = make_blank_pdf_bytes()
    files = {"file": ("scanned_blank.pdf", blank_pdf_bytes, "application/pdf")}
    response = client.post("/api/resumes/upload", files=files)

    assert response.status_code == 200, f"Expected 200, got {response.status_code}: {response.text}"
    data = response.json()
    assert data["parsing_status"] == "FAILED"
    print("  [PASS] test_scanned_or_blank_pdf_graceful_handling: Blank PDF returned FAILED status gracefully.")


def test_filename_sanitization(client):
    """TC-4.1-09: Filenames with path traversal characters are properly sanitized."""
    reset_rate_limiter()
    pdf_bytes = make_valid_pdf_bytes("Alice Smith - DevOps Engineer")
    files = {"file": ("../../etc/passwd_evil.pdf", pdf_bytes, "application/pdf")}
    response = client.post("/api/resumes/upload", files=files)

    assert response.status_code == 200, f"Expected 200, got {response.status_code}"
    data = response.json()
    assert "/" not in data["file_name"]
    assert ".." not in data["file_name"]
    print(f"  [PASS] test_filename_sanitization: Sanitized to '{data['file_name']}'.")


def test_oversized_file_rejected(client):
    """TC-4.1-10: File exceeding 5 MB limit rejected with 400 Bad Request."""
    # Reset rate limit to ensure test isolation
    import backend.services.rate_limiter as rl
    rl._UPLOAD_HISTORY.clear()

    oversized_bytes = b"%PDF-1.4\n" + b"X" * (5 * 1024 * 1024 + 1024)
    files = {"file": ("large_resume.pdf", oversized_bytes, "application/pdf")}
    response = client.post("/api/resumes/upload", files=files)

    assert response.status_code == 400, f"Expected 400, got {response.status_code}"
    assert "5 mb limit" in response.json()["detail"].lower()
    print("  [PASS] test_oversized_file_rejected: File exceeding 5 MB rejected with 400.")


def run_all_tests():
    print("\n=== Running Stage 4.1 Acceptance Tests (Resume Upload & Text Extraction) ===")
    setup_sqlite_test_env(app)
    client = TestClient(app)

    pdf_id, pdf_token = test_valid_pdf_upload(client)
    docx_id, docx_token = test_valid_docx_upload(client)
    test_get_resume_metadata_endpoint(client, pdf_id, pdf_token)
    test_get_nonexistent_resume_metadata(client)
    test_invalid_extension_rejected(client)
    test_spoofed_signature_rejected(client)
    test_empty_zero_byte_file_rejected(client)
    test_scanned_or_blank_pdf_graceful_handling(client)
    test_filename_sanitization(client)
    test_oversized_file_rejected(client)

    print("\n=== All Stage 4.1 Acceptance Tests Passed Successfully! ===\n")


if __name__ == "__main__":
    run_all_tests()
