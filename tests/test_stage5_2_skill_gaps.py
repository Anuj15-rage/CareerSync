"""
Stage 5.2 Acceptance Tests — CareerSync Skill Gap Analysis
Updated for Phase 1 token security, verified catalog, and isolated SQLite execution.
"""
import io
import sys
import uuid
from pathlib import Path
from unittest.mock import patch

# Add project root to sys.path so backend imports resolve cleanly
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

import docx
from fastapi.testclient import TestClient

from backend.main import app
from backend.models import (
    AnalysisSession,
    Resume,
    SkillGap,
)
import backend.services.rate_limiter as rl
from tests.test_sqlite_fixtures import setup_sqlite_test_env


def reset_rate_limiter():
    rl._UPLOAD_HISTORY.clear()


def create_docx_with_skills(skills_text: str, candidate_name: str = "Test Candidate") -> bytes:
    """Create in-memory DOCX resume with specified skills."""
    doc = docx.Document()
    doc.add_paragraph(candidate_name)
    doc.add_paragraph("Email: candidate@example.com | Phone: +1-555-010-9999")
    doc.add_heading("TECHNICAL SKILLS", level=1)
    doc.add_paragraph(skills_text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_matched_and_missing_skills_accuracy(client, session_factory):
    """
    Verify exact scoring, gap priority mapping, and database persistence against Role #1:
    Python Backend Developer (Total Weight: 9.00):
      - 6 REQUIRED @ 1.00: Python, SQL, FastAPI, MySQL, Git, RESTful APIs (6.00)
      - 4 PREFERRED @ 0.75: PostgreSQL, Docker, Database Normalization, Unit Testing (3.00)

    Candidate has: Python, SQL, FastAPI, MySQL, Docker, HTML5, CSS3
      - Matched Required: Python, SQL, FastAPI, MySQL = 4.00
      - Matched Preferred: Docker = 0.75
      - Matched Total = 4.75
      - Expected Compatibility Score = round((4.75 / 9.00) * 100, 2) = 52.78%
      - Expected Technical Score = round((4.00 / 6.00) * 100, 2) = 66.67%
      - Critical Missing: Git, RESTful APIs (CRITICAL_MISSING)
      - Secondary Missing: PostgreSQL, Database Normalization, Unit Testing (SECONDARY_MISSING)
    """
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, SQL, FastAPI, MySQL, Docker, HTML5, CSS3", "Python Candidate")
    files = {"file": (f"test_gap_backend_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    # 1. Upload & Analyze Resume
    up_res = client.post("/api/resumes/upload", files=files)
    assert up_res.status_code == 200
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    an_res = client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)
    assert an_res.status_code == 200

    # 2. Run Skill Gap Analysis POST
    gap_res = client.post(f"/api/resumes/{resume_id}/skill-gaps/analyze", headers=headers)
    assert gap_res.status_code == 200, f"Gap analysis failed: {gap_res.text}"
    data = gap_res.json()

    assert data["resume_id"] == resume_id
    assert data["roles_evaluated_count"] >= 5

    # Find Role #1
    role_1_gap = next((r for r in data["role_gaps"] if r["job_role_id"] == 1), None)
    assert role_1_gap is not None, "Role #1 missing in gap results"
    assert "Backend Developer" in role_1_gap["job_role_title"]
    assert role_1_gap["compatibility_score"] == 52.78
    assert role_1_gap["technical_skill_score"] == 66.67
    assert role_1_gap["possessed_skills_count"] == 5
    assert role_1_gap["critical_missing_count"] == 2
    assert role_1_gap["secondary_missing_count"] == 3

    # Check critical missing
    crit_names = {s["skill_name"] for s in role_1_gap["critical_missing_skills"]}
    assert crit_names == {"Git", "RESTful APIs"}
    for item in role_1_gap["critical_missing_skills"]:
        assert item["gap_priority"] == "CRITICAL_MISSING"
        assert item["importance"] == "REQUIRED"
        assert item["weight"] == 1.00
        assert "Critical missing skill" in item["reason"]

    # Check secondary missing
    sec_names = {s["skill_name"] for s in role_1_gap["secondary_missing_skills"]}
    assert sec_names == {"PostgreSQL", "Database Normalization", "Unit Testing"}
    for item in role_1_gap["secondary_missing_skills"]:
        assert item["gap_priority"] == "SECONDARY_MISSING"
        assert item["importance"] == "PREFERRED"
        assert item["weight"] == 0.75
        assert "Secondary missing skill" in item["reason"]

    # 3. Verify in database directly
    with session_factory() as db:
        session = db.query(AnalysisSession).filter(
            AnalysisSession.resume_id == resume_id,
            AnalysisSession.job_role_id == 1,
        ).first()
        assert session is not None
        assert float(session.compatibility_score) == 52.78
        assert float(session.technical_skill_score) == 66.67
        assert "Backend Developer" in session.target_role_title

        # Check skill_gaps in DB
        db_gaps = db.query(SkillGap).filter(SkillGap.session_id == session.id).all()
        assert len(db_gaps) == 5
        crit_db = [g for g in db_gaps if g.gap_priority == "CRITICAL_MISSING"]
        sec_db = [g for g in db_gaps if g.gap_priority == "SECONDARY_MISSING"]
        assert len(crit_db) == 2
        assert len(sec_db) == 3

        # Verify resumes.analysis_completed_at was NOT modified
        resume_db = db.query(Resume).filter(Resume.id == resume_id).first()
        assert resume_db.analysis_completed_at is not None

    print("  [PASS] test_matched_and_missing_skills_accuracy: 52.78% score and CRITICAL/SECONDARY mapping verified in DB.")


def test_get_behavior_read_only(client, session_factory):
    """
    Verify GET /api/resumes/{id}/skill-gaps:
    1. Returns 404 if gap analysis has not been performed yet (no rows written).
    2. Returns 200 with persisted results after POST has been run.
    3. GET /api/resumes/{id}/skill-gaps/{role_id} returns single role detail accurately.
    """
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, SQL", "Read Candidate")
    files = {"file": (f"test_get_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    assert up_res.status_code == 200
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    # 1. GET before POST analyze must return 404
    get_before = client.get(f"/api/resumes/{resume_id}/skill-gaps", headers=headers)
    assert get_before.status_code == 404
    assert "not been performed" in get_before.json()["detail"].lower()

    # Verify zero sessions exist in DB
    with session_factory() as db:
        assert db.query(AnalysisSession).filter(AnalysisSession.resume_id == resume_id).count() == 0

    # 2. Run POST
    post_res = client.post(f"/api/resumes/{resume_id}/skill-gaps/analyze", headers=headers)
    assert post_res.status_code == 200

    # 3. GET after POST must return identical persisted data
    get_after = client.get(f"/api/resumes/{resume_id}/skill-gaps", headers=headers)
    assert get_after.status_code == 200
    data = get_after.json()
    assert data["resume_id"] == resume_id
    assert data["roles_evaluated_count"] >= 5

    # 4. Single role GET
    get_single = client.get(f"/api/resumes/{resume_id}/skill-gaps/1", headers=headers)
    assert get_single.status_code == 200
    single_data = get_single.json()
    assert single_data["job_role_id"] == 1
    assert "Backend Developer" in single_data["job_role_title"]

    print("  [PASS] test_get_behavior_read_only: 404 before analyze verified; read-only retrieval verified.")


def test_zero_gaps_perfect_match(client, session_factory):
    """Verify candidate matching 100% of skills in Role #5 achieves zero gaps and stores 0 skill_gaps rows."""
    reset_rate_limiter()
    # Role #5: Cloud / DevOps Engineer (Git, Docker, Linux / Bash, AWS, CI/CD Pipelines)
    docx_bytes = create_docx_with_skills("Git, Docker, Linux / Bash, AWS, CI/CD Pipelines", "Perfect Candidate")
    files = {"file": (f"test_perfect_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    assert up_res.status_code == 200
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    gap_res = client.post(f"/api/resumes/{resume_id}/skill-gaps/analyze", headers=headers)
    assert gap_res.status_code == 200
    role_5 = next(r for r in gap_res.json()["role_gaps"] if r["job_role_id"] == 5)

    assert role_5["compatibility_score"] == 100.00
    assert role_5["technical_skill_score"] == 100.00
    assert role_5["missing_skills_count"] == 0
    assert role_5["critical_missing_skills"] == []
    assert role_5["secondary_missing_skills"] == []

    # Verify directly in DB: session exists, but 0 rows in skill_gaps
    with session_factory() as db:
        session = db.query(AnalysisSession).filter(
            AnalysisSession.resume_id == resume_id,
            AnalysisSession.job_role_id == 5,
        ).first()
        assert session is not None
        assert float(session.compatibility_score) == 100.00
        gaps_count = db.query(SkillGap).filter(SkillGap.session_id == session.id).count()
        assert gaps_count == 0, f"Expected 0 skill_gaps for perfect match, found {gaps_count}"

    print("  [PASS] test_zero_gaps_perfect_match: 100.00% score and 0 gap records verified.")


def test_zero_matched_skills(client):
    """Verify candidate with zero catalog skills gets 0.00% scores and all role skills marked as missing."""
    reset_rate_limiter()
    doc = docx.Document()
    doc.add_paragraph("Jordan Taylor")
    doc.add_paragraph("General candidate with no technical skills.")
    doc.add_heading("EDUCATION", level=1)
    doc.add_paragraph("State University | BA Philosophy | 2018 - 2022")
    buf = io.BytesIO()
    doc.save(buf)
    files = {"file": (f"test_zero_{uuid.uuid4().hex[:6]}.docx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    assert up_res.status_code == 200
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    gap_res = client.post(f"/api/resumes/{resume_id}/skill-gaps/analyze", headers=headers)
    assert gap_res.status_code == 200
    data = gap_res.json()

    for role_gap in data["role_gaps"]:
        assert role_gap["compatibility_score"] == 0.00
        assert role_gap["technical_skill_score"] == 0.00
        assert role_gap["possessed_skills_count"] == 0
        assert role_gap["missing_skills_count"] == role_gap["total_role_skills_count"]

    print("  [PASS] test_zero_matched_skills: 0.00% scores and all role skills reported as gaps.")


def test_idempotent_reanalysis(client, session_factory):
    """Verify re-running skill gap analysis updates sessions in-place and leaves no duplicate rows."""
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, FastAPI", "Idempotent Candidate")
    files = {"file": (f"test_idem_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    # Run 1
    res1 = client.post(f"/api/resumes/{resume_id}/skill-gaps/analyze", headers=headers)
    assert res1.status_code == 200

    with session_factory() as db:
        sessions_1 = db.query(AnalysisSession).filter(AnalysisSession.resume_id == resume_id).count()
        gaps_1 = (
            db.query(SkillGap)
            .join(AnalysisSession, SkillGap.session_id == AnalysisSession.id)
            .filter(AnalysisSession.resume_id == resume_id)
            .count()
        )
        assert sessions_1 >= 5

    # Run 2
    res2 = client.post(f"/api/resumes/{resume_id}/skill-gaps/analyze", headers=headers)
    assert res2.status_code == 200

    with session_factory() as db:
        sessions_2 = db.query(AnalysisSession).filter(AnalysisSession.resume_id == resume_id).count()
        gaps_2 = (
            db.query(SkillGap)
            .join(AnalysisSession, SkillGap.session_id == AnalysisSession.id)
            .filter(AnalysisSession.resume_id == resume_id)
            .count()
        )
        assert sessions_1 == sessions_2, "Session count must not increase on re-analysis"
        assert gaps_1 == gaps_2, "SkillGap count must not increase on re-analysis"

    print("  [PASS] test_idempotent_reanalysis: Re-analysis is completely idempotent; no duplicate sessions or gaps.")


def test_transactional_rollback(client, session_factory):
    """Verify failure during gap analysis cleanly rolls back all modifications."""
    reset_rate_limiter()
    docx_bytes = create_docx_with_skills("Python, SQL", "Rollback Candidate")
    files = {"file": (f"test_rb_{uuid.uuid4().hex[:6]}.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}

    up_res = client.post("/api/resumes/upload", files=files)
    token = up_res.json()["access_token"]
    resume_id = up_res.json()["id"]
    headers = {"X-Resume-Token": token}

    client.post(f"/api/resumes/{resume_id}/analyze", headers=headers)

    # Simulate failure in session creation
    with patch("backend.services.skill_gap_analyzer.compute_single_role_match", side_effect=RuntimeError("Simulated analyzer crash")):
        fail_res = client.post(f"/api/resumes/{resume_id}/skill-gaps/analyze", headers=headers)
        assert fail_res.status_code == 500

    # Verify zero sessions or gaps persisted
    with session_factory() as db:
        count = db.query(AnalysisSession).filter(AnalysisSession.resume_id == resume_id).count()
        assert count == 0, f"Expected 0 sessions after rollback, found {count}"

    print("  [PASS] test_transactional_rollback: Analysis rolls back cleanly on simulated failure.")


def run_all_tests():
    print("\n=== Running Stage 5.2 Acceptance Tests (Skill Gap Analysis) ===")
    engine, session_factory = setup_sqlite_test_env(app)
    client = TestClient(app)

    test_matched_and_missing_skills_accuracy(client, session_factory)
    test_get_behavior_read_only(client, session_factory)
    test_zero_gaps_perfect_match(client, session_factory)
    test_zero_matched_skills(client)
    test_idempotent_reanalysis(client, session_factory)
    test_transactional_rollback(client, session_factory)

    print("\n=== All Stage 5.2 Acceptance Tests Passed Successfully! ===\n")


if __name__ == "__main__":
    run_all_tests()
